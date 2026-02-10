import streamlit as st
import os
import json
import datetime
from dotenv import load_dotenv
from groq import Groq
from presidio_analyzer import AnalyzerEngine, PatternRecognizer, Pattern
from presidio_anonymizer import AnonymizerEngine
import pypdf
import docx
import pandas as pd

# CONFIGURATION
load_dotenv()
GROQ_API_KEY = os.getenv("GROQ_API_KEY")
GROQ_MODEL = "llama-3.3-70b-versatile"
AUDIT_FILE = "audit_log.json"

@st.cache_resource
def load_tools():
    analyzer = AnalyzerEngine()
    anonymizer = AnonymizerEngine()

    # Custom Detectors
    ssn_pattern = Pattern(name="ssn_pattern", regex=r"\b\d{3}-\d{2}-\d{4}\b", score=0.9)
    ssn_recognizer = PatternRecognizer(supported_entity="US_SSN", name="Force_SSN", patterns=[ssn_pattern])
    analyzer.registry.add_recognizer(ssn_recognizer)

    cc_pattern = Pattern(name="cc_pattern", regex=r"\b\d{4}-\d{4}-\d{4}-\d{4}\b", score=0.9)
    cc_recognizer = PatternRecognizer(supported_entity="CREDIT_CARD", name="Force_CC", patterns=[cc_pattern])
    analyzer.registry.add_recognizer(cc_recognizer)

    return analyzer, anonymizer

analyzer, anonymizer = load_tools()

# JARGON 
def add_jargon_recognizer(analyzer_engine, jargon_list):
    if not jargon_list: return
    jargon_recognizer = PatternRecognizer(supported_entity="CUSTOM_JARGON", name="Jargon_List", deny_list=jargon_list)
    try: analyzer_engine.registry.remove_recognizer("Jargon_List")
    except: pass
    analyzer_engine.registry.add_recognizer(jargon_recognizer)

def validate_input(text):
    if not text: return None, "⚠️ Input is empty."
    clean_text = text.strip()
    if len(clean_text) > 10000: return None, "⚠️ Input too long."
    forbidden = ["ignore previous instructions", "system override"]
    for phrase in forbidden:
        if phrase in clean_text.lower(): return None, "Prompt Injection Detected."
    return clean_text, None

# AUDIT LOGGER 
def log_audit_event(original_len, secret_map):
    pii_counts = {}
    for item in secret_map:
        pii_counts[item.entity_type] = pii_counts.get(item.entity_type, 0) + 1
    
    log_entry = {
        "timestamp": datetime.datetime.now().isoformat(),
        "event": "DATA_MASKING",
        "input_length": original_len,
        "blocked_items": len(secret_map),
        "risk_types": list(pii_counts.keys()),
        "details": str(pii_counts)
    }
    
    try:
        if not os.path.exists(AUDIT_FILE):
            with open(AUDIT_FILE, "w") as f: json.dump([], f)
        with open(AUDIT_FILE, "r+") as f:
            try: data = json.load(f)
            except: data = []
            data.append(log_entry)
            f.seek(0)
            json.dump(data, f, indent=4)
    except Exception: pass

def read_file(uploaded_file):
    text = ""
    try:
        if uploaded_file.name.endswith(".pdf"):
            reader = pypdf.PdfReader(uploaded_file)
            for page in reader.pages: text += page.extract_text() + "\n"
        elif uploaded_file.name.endswith(".docx"):
            doc = docx.Document(uploaded_file)
            for para in doc.paragraphs: text += para.text + "\n"
        else: text = uploaded_file.read().decode("utf-8")
    except Exception as e: return f"Error: {e}"
    return text

def mask_pii(text):
    target_entities = ["PERSON", "EMAIL_ADDRESS", "PHONE_NUMBER", "LOCATION", "CREDIT_CARD", "US_SSN", "US_PASSPORT", "IBAN_CODE", "CUSTOM_JARGON"]
    results = analyzer.analyze(text=text, language='en', entities=target_entities, score_threshold=0.4)
    anonymized_result = anonymizer.anonymize(text=text, analyzer_results=results)
    log_audit_event(len(text), anonymized_result.items)
    return anonymized_result.text, anonymized_result.items

def unmask_pii(ai_response, items, original_text):
    processed_response = ai_response
    for item in items:
        placeholder = f"<{item.entity_type}>"
        real_value = original_text[item.start:item.end]
        processed_response = processed_response.replace(placeholder, real_value)
    return processed_response

def ask_groq(safe_text):
    if not GROQ_API_KEY: return "Key missing in .env"
    try:
        client = Groq(api_key=GROQ_API_KEY)
        chat_completion = client.chat.completions.create(
            messages=[
                {"role": "system", "content": "You are a helpful assistant. Preserve placeholders like <PERSON> exactly."},
                {"role": "user", "content": safe_text}
            ],
            model=GROQ_MODEL, temperature=0.7,
        )
        return chat_completion.choices[0].message.content
    except Exception as e: return f"Cloud Error: {str(e)}"

def load_premium_css():
    st.markdown("""
    <style>
    /* Import Premium Fonts & Icons */
    @import url('https://fonts.googleapis.com/css2?family=Inter:wght@300;400;500;600;700;800&family=JetBrains+Mono:wght@400;500;600&display=swap');
    @import url('https://cdnjs.cloudflare.com/ajax/libs/font-awesome/6.4.0/css/all.min.css');
    
    /* Global Resets & Base */
    * {
        font-family: 'Inter', -apple-system, BlinkMacSystemFont, sans-serif;
    }
    
    /* Main Container - Pure White Sophistication */
    [data-testid="stAppViewContainer"] {
        background: #ffffff;
        background-image: 
            radial-gradient(at 0% 0%, rgba(99, 102, 241, 0.03) 0px, transparent 50%),
            radial-gradient(at 100% 0%, rgba(59, 130, 246, 0.03) 0px, transparent 50%),
            radial-gradient(at 100% 100%, rgba(139, 92, 246, 0.03) 0px, transparent 50%);
    }
    
    [data-testid="stHeader"] {
        background: transparent;
    }
    
    /* Sidebar - Minimalist White Panel */
    [data-testid="stSidebar"] {
        background: linear-gradient(180deg, #fafafa 0%, #ffffff 100%);
        border-right: 1px solid #e5e7eb;
        box-shadow: 2px 0 24px rgba(0, 0, 0, 0.04);
    }
    
    [data-testid="stSidebar"] [data-testid="stMarkdownContainer"] h2 {
        color: #111827;
        font-weight: 600;
        font-size: 0.9rem;
        letter-spacing: 0.5px;
        text-transform: uppercase;
        margin-bottom: 1.5rem;
        padding-bottom: 0.8rem;
        border-bottom: 2px solid #e5e7eb;
    }
    
    /* Main Title - Editorial Sophistication */
    h1 {
        color: #111827;
        font-weight: 800;
        font-size: 2.75rem !important;
        letter-spacing: -1.5px;
        margin-bottom: 0.5rem;
        animation: fadeInDown 0.8s ease-out;
        line-height: 1.1;
    }
    
    @keyframes fadeInDown {
        from { opacity: 0; transform: translateY(-20px); }
        to { opacity: 1; transform: translateY(0); }
    }
    
    /* Tab System - Clean & Refined */
    .stTabs [data-baseweb="tab-list"] {
        gap: 0;
        background: #f9fafb;
        padding: 0.375rem;
        border-radius: 14px;
        border: 1px solid #e5e7eb;
        box-shadow: inset 0 1px 2px rgba(0, 0, 0, 0.04);
    }
    
    .stTabs [data-baseweb="tab"] {
        background: transparent;
        border-radius: 10px;
        padding: 0.625rem 1.5rem;
        color: #6b7280;
        font-weight: 500;
        font-size: 0.9rem;
        transition: all 0.25s cubic-bezier(0.4, 0, 0.2, 1);
        border: none;
    }
    
    .stTabs [data-baseweb="tab"]:hover {
        background: rgba(99, 102, 241, 0.06);
        color: #4f46e5;
    }
    
    .stTabs [aria-selected="true"] {
        background: #ffffff;
        color: #4f46e5 !important;
        box-shadow: 
            0 1px 3px rgba(0, 0, 0, 0.08),
            0 0 0 1px rgba(99, 102, 241, 0.1);
        font-weight: 600;
    }
    
    /* Premium Card System */
    [data-testid="column"] {
        background: #ffffff;
        border-radius: 16px;
        padding: 2rem;
        border: 1px solid #e5e7eb;
        box-shadow: 
            0 1px 3px rgba(0, 0, 0, 0.04),
            0 20px 25px -5px rgba(0, 0, 0, 0.02);
        transition: all 0.35s cubic-bezier(0.4, 0, 0.2, 1);
    }
    
    [data-testid="column"]:hover {
        transform: translateY(-2px);
        box-shadow: 
            0 4px 6px -1px rgba(0, 0, 0, 0.06),
            0 20px 25px -5px rgba(99, 102, 241, 0.08);
        border-color: rgba(99, 102, 241, 0.2);
    }
    
    /* Input Fields - Crisp & Professional */
    .stTextArea textarea, .stTextInput input {
        background: #f9fafb !important;
        border: 1.5px solid #e5e7eb !important;
        border-radius: 12px !important;
        color: #111827 !important;
        padding: 1rem !important;
        font-size: 0.9rem !important;
        transition: all 0.25s ease;
        line-height: 1.6;
    }
    
    .stTextArea textarea:focus, .stTextInput input:focus {
        border-color: #6366f1 !important;
        box-shadow: 0 0 0 3px rgba(99, 102, 241, 0.1) !important;
        background: #ffffff !important;
    }
    
    .stTextArea textarea::placeholder, .stTextInput input::placeholder {
        color: #9ca3af !important;
    }
    
    /* Buttons - Sophisticated Depth */
    .stButton button {
        background: linear-gradient(135deg, #4f46e5 0%, #6366f1 100%);
        color: white;
        border: none;
        border-radius: 12px;
        padding: 0.875rem 2rem;
        font-weight: 600;
        font-size: 0.875rem;
        letter-spacing: 0.3px;
        box-shadow: 
            0 4px 6px -1px rgba(79, 70, 229, 0.2),
            0 2px 4px -1px rgba(79, 70, 229, 0.1);
        transition: all 0.25s cubic-bezier(0.4, 0, 0.2, 1);
        text-transform: uppercase;
        position: relative;
    }
    
    .stButton button:hover {
        transform: translateY(-1px);
        box-shadow: 
            0 10px 15px -3px rgba(79, 70, 229, 0.25),
            0 4px 6px -2px rgba(79, 70, 229, 0.15);
        background: linear-gradient(135deg, #4338ca 0%, #4f46e5 100%);
    }
    
    .stButton button:active {
        transform: translateY(0);
    }
    
    /* Download Button - Success Variant */
    .stDownloadButton button {
        background: linear-gradient(135deg, #059669 0%, #10b981 100%);
        box-shadow: 
            0 4px 6px -1px rgba(5, 150, 105, 0.2),
            0 2px 4px -1px rgba(5, 150, 105, 0.1);
    }
    
    .stDownloadButton button:hover {
        background: linear-gradient(135deg, #047857 0%, #059669 100%);
        box-shadow: 
            0 10px 15px -3px rgba(5, 150, 105, 0.25),
            0 4px 6px -2px rgba(5, 150, 105, 0.15);
    }
    
    /* Metrics - Executive Dashboard */
    [data-testid="stMetricValue"] {
        color: #111827;
        font-size: 2.5rem !important;
        font-weight: 700;
        letter-spacing: -1px;
    }
    
    [data-testid="stMetricLabel"] {
        color: #6b7280;
        font-size: 0.8rem;
        font-weight: 600;
        text-transform: uppercase;
        letter-spacing: 0.8px;
    }
    
    [data-testid="stMetricDelta"] {
        font-size: 0.85rem;
        font-weight: 500;
    }
    
    /* Status Messages - Refined Alerts */
    .stSuccess, .stError, .stInfo, .stWarning {
        background: #ffffff !important;
        backdrop-filter: none;
        border-radius: 12px !important;
        border: 1px solid;
        border-left: 4px solid;
        padding: 1rem 1.25rem !important;
        animation: slideInRight 0.4s ease-out;
        box-shadow: 0 1px 3px rgba(0, 0, 0, 0.06);
    }
    
    @keyframes slideInRight {
        from { opacity: 0; transform: translateX(20px); }
        to { opacity: 1; transform: translateX(0); }
    }
    
    .stSuccess {
        border-color: #d1fae5 !important;
        border-left-color: #10b981 !important;
        color: #065f46 !important;
        background: #f0fdf4 !important;
    }
    
    .stError {
        border-color: #fecaca !important;
        border-left-color: #ef4444 !important;
        color: #991b1b !important;
        background: #fef2f2 !important;
    }
    
    .stInfo {
        border-color: #dbeafe !important;
        border-left-color: #3b82f6 !important;
        color: #1e40af !important;
        background: #eff6ff !important;
    }
    
    .stWarning {
        border-color: #fef3c7 !important;
        border-left-color: #f59e0b !important;
        color: #92400e !important;
        background: #fffbeb !important;
    }
    
    /* Expander - Minimal Elegance */
    .streamlit-expanderHeader {
        background: #f9fafb;
        border-radius: 10px;
        padding: 0.875rem 1rem;
        color: #374151;
        font-weight: 500;
        border: 1px solid #e5e7eb;
        transition: all 0.25s ease;
    }
    
    .streamlit-expanderHeader:hover {
        background: #f3f4f6;
        border-color: #d1d5db;
    }
    
    /* Code Blocks - Editor Style */
    code {
        background: #f9fafb !important;
        border: 1px solid #e5e7eb;
        border-radius: 6px;
        padding: 0.2rem 0.5rem;
        color: #4f46e5;
        font-family: 'JetBrains Mono', monospace;
        font-size: 0.85rem;
    }
    
    pre {
        background: #fafafa !important;
        border: 1px solid #e5e7eb;
        border-radius: 12px;
        padding: 1.25rem;
    }
    
    pre code {
        background: transparent !important;
        border: none;
        padding: 0;
        color: #111827;
    }
    
    /* Radio Buttons - Clean Toggle */
    .stRadio > label {
        color: #374151;
        font-weight: 500;
        font-size: 0.9rem;
    }
    
    .stRadio [role="radiogroup"] {
        background: #f9fafb;
        padding: 0.5rem;
        border-radius: 12px;
        gap: 0.5rem;
        border: 1px solid #e5e7eb;
    }
    
    /* File Uploader - Dashed Border Style */
    [data-testid="stFileUploader"] {
        background: #fafafa;
        border: 2px dashed #d1d5db;
        border-radius: 12px;
        padding: 2rem;
        transition: all 0.25s ease;
    }
    
    [data-testid="stFileUploader"]:hover {
        border-color: #6366f1;
        background: #f9fafb;
    }
    
    /* DataFrames - Clean Tables */
    [data-testid="stDataFrame"] {
        border-radius: 12px;
        overflow: hidden;
        border: 1px solid #e5e7eb;
    }
    
    /* Charts - Minimal Containers */
    [data-testid="stArrowVegaLiteChart"], 
    [data-testid="stLineChart"],
    [data-testid="stBarChart"] {
        background: #fafafa;
        border-radius: 12px;
        padding: 1.25rem;
        border: 1px solid #e5e7eb;
    }
    
    /* Dividers */
    hr {
        border-color: #e5e7eb;
        margin: 2.5rem 0;
    }
    
    /* Subheaders */
    h2, h3, h4 {
        color: #111827;
        font-weight: 700;
        letter-spacing: -0.5px;
        margin-top: 1.5rem;
        margin-bottom: 1rem;
    }
    
    h2 {
        font-size: 1.875rem;
    }
    
    h3 {
        font-size: 1.5rem;
    }
    
    h4 {
        font-size: 1.125rem;
    }
    
    /* Spinner/Loading States */
    .stSpinner > div {
        border-color: #6366f1 transparent transparent transparent !important;
    }
    
    /* Scrollbar Styling */
    ::-webkit-scrollbar {
        width: 10px;
        height: 10px;
    }
    
    ::-webkit-scrollbar-track {
        background: #f3f4f6;
    }
    
    ::-webkit-scrollbar-thumb {
        background: #d1d5db;
        border-radius: 5px;
    }
    
    ::-webkit-scrollbar-thumb:hover {
        background: #9ca3af;
    }
    
    /* Status Widget */
    [data-testid="stStatusWidget"] {
        background: #ffffff;
        border-radius: 12px;
        border: 1px solid #e5e7eb;
        box-shadow: 0 1px 3px rgba(0, 0, 0, 0.06);
    }
    
    /* Labels */
    label {
        color: #374151 !important;
        font-weight: 500 !important;
        font-size: 0.875rem !important;
    }
    
    /* Custom Alert Boxes */
    .custom-alert {
        padding: 1rem 1.25rem;
        border-radius: 12px;
        border: 1px solid;
        margin: 1rem 0;
        display: flex;
        align-items: center;
        gap: 0.75rem;
        font-size: 0.9rem;
        line-height: 1.6;
    }
    
    .alert-success {
        background: #f0fdf4;
        border-color: #d1fae5;
        border-left: 4px solid #10b981;
        color: #065f46;
    }
    
    .alert-info {
        background: #eff6ff;
        border-color: #dbeafe;
        border-left: 4px solid #3b82f6;
        color: #1e40af;
    }
    
    .alert-warning {
        background: #fffbeb;
        border-color: #fef3c7;
        border-left: 4px solid #f59e0b;
        color: #92400e;
    }
    
    .alert-error {
        background: #fef2f2;
        border-color: #fecaca;
        border-left: 4px solid #ef4444;
        color: #991b1b;
    }
    
    /* Section Headers with Icons */
    .section-header {
        display: flex;
        align-items: center;
        gap: 0.75rem;
        margin-bottom: 1.5rem;
        padding-bottom: 1rem;
        border-bottom: 2px solid #f3f4f6;
    }
    
    .section-header i {
        font-size: 1.25rem;
        color: #6366f1;
    }
    
    .section-header h3 {
        margin: 0;
        font-size: 1.25rem;
        font-weight: 700;
        color: #111827;
    }
    
    /* Metric Cards */
    .metric-card {
        background: linear-gradient(135deg, #f9fafb 0%, #ffffff 100%);
        border: 1px solid #e5e7eb;
        border-radius: 14px;
        padding: 1.5rem;
        text-align: center;
        transition: all 0.3s ease;
    }
    
    .metric-card:hover {
        transform: translateY(-2px);
        box-shadow: 0 4px 12px rgba(0, 0, 0, 0.06);
        border-color: #d1d5db;
    }
    
    .metric-card i {
        font-size: 2rem;
        margin-bottom: 0.75rem;
        display: block;
    }
    </style>
    """, unsafe_allow_html=True)

st.set_page_config(
    page_title="Enterprise Privacy Gateway", 
    page_icon="🛡️",
    layout="wide",
    initial_sidebar_state="expanded"
)

load_premium_css()

st.markdown("""
    <div style='text-align: center; margin-bottom: 3.5rem; padding: 2rem 0;'>
        <div style='margin-bottom: 1.5rem;'>
            <svg width="72" height="72" viewBox="0 0 24 24" fill="none" xmlns="http://www.w3.org/2000/svg" style="filter: drop-shadow(0 4px 12px rgba(99, 102, 241, 0.2));">
                <path d="M12 2L3 7V11C3 16.55 6.84 21.74 12 23C17.16 21.74 21 16.55 21 11V7L12 2Z" fill="url(#gradient1)"/>
                <path d="M12 8C10.9 8 10 8.9 10 10V14C10 15.1 10.9 16 12 16C13.1 16 14 15.1 14 14V10C14 8.9 13.1 8 12 8Z" fill="white"/>
                <defs>
                    <linearGradient id="gradient1" x1="3" y1="2" x2="21" y2="23" gradientUnits="userSpaceOnUse">
                        <stop stop-color="#6366f1"/>
                        <stop offset="1" stop-color="#4f46e5"/>
                    </linearGradient>
                </defs>
            </svg>
        </div>
        <h1 style='margin-bottom: 0.75rem;'>Enterprise Privacy Gateway</h1>
        <p style='color: #6b7280; font-size: 1rem; font-weight: 500; max-width: 600px; margin: 0 auto; line-height: 1.6;'>
            Military-grade data protection and compliance infrastructure for modern enterprises
        </p>
    </div>
""", unsafe_allow_html=True)

with st.sidebar:
    st.markdown("""
        <div style='margin-bottom: 2rem; padding-bottom: 1rem; border-bottom: 2px solid #e5e7eb;'>
            <div style='display: flex; align-items: center; gap: 0.75rem;'>
                <i class="fas fa-filter" style='color: #6366f1; font-size: 1.1rem;'></i>
                <h3 style='margin: 0; font-size: 0.9rem; font-weight: 700; color: #111827; text-transform: uppercase; letter-spacing: 0.5px;'>Custom Blocklist</h3>
            </div>
            <p style='margin-top: 0.5rem; font-size: 0.8rem; color: #6b7280; line-height: 1.5;'>Define sensitive terms to protect</p>
        </div>
    """, unsafe_allow_html=True)
    
    jargon_input = st.text_area(
        "Protected Terms:", 
        value="Project Apollo, Skynet",
        help="Comma-separated list of confidential terms to mask",
        height=120
    )
    jargon_list = [word.strip() for word in jargon_input.split(",") if word.strip()]
    if jargon_list:
        add_jargon_recognizer(analyzer, jargon_list)
        st.markdown(f"""
            <div class='custom-alert alert-success' style='margin-top: 1rem;'>
                <i class="fas fa-check-circle" style='color: #10b981;'></i>
                <span><strong>{len(jargon_list)}</strong> terms actively protected</span>
            </div>
        """, unsafe_allow_html=True)

tab_user, tab_auditor = st.tabs([
    "Secure Workspace",
    "Analytics Dashboard"
])

with tab_user:
    col1, col2 = st.columns(2, gap="large")
    
    with col1:
        st.markdown("""
            <div class='section-header'>
                <i class="fas fa-pen-to-square"></i>
                <h3>Input Zone</h3>
            </div>
        """, unsafe_allow_html=True)
        
        input_type = st.radio(
            "Data Source:", 
            ["Text", "File Upload"], 
            horizontal=True,
            label_visibility="collapsed"
        )
        
        user_input = ""
        if input_type == "Text":
            user_input = st.text_area(
                "Input:", 
                height=280, 
                placeholder="Paste sensitive content here... (emails, documents, reports)",
                label_visibility="collapsed"
            )
        else:
            uploaded = st.file_uploader(
                "Upload Document", 
                type=["txt", "pdf", "docx"],
                label_visibility="collapsed"
            )
            if uploaded: 
                user_input = read_file(uploaded)
                st.markdown(f"""
                    <div class='custom-alert alert-info'>
                        <i class="fas fa-file-check"></i>
                        <span>Document loaded: <strong>{uploaded.name}</strong></span>
                    </div>
                """, unsafe_allow_html=True)

        st.markdown("<br>", unsafe_allow_html=True)
        if st.button("SECURE & PROCESS", use_container_width=True, type="primary"):
            valid_text, error = validate_input(user_input)
            if error: 
                st.error(error)
            else:
                safe_text, secret_map = mask_pii(valid_text)
                with st.status("Processing through privacy layer...", expanded=True) as status:
                    ai_answer = ask_groq(safe_text)
                    if "🚨" in ai_answer:
                        st.error(ai_answer)
                        status.update(label="Processing Failed", state="error")
                    else:
                        final_answer = unmask_pii(ai_answer, secret_map, valid_text)
                        status.update(label="Complete", state="complete")
                        with col2:
                            st.markdown("""
                                <div class='section-header'>
                                    <i class="fas fa-shield-check" style='color: #10b981;'></i>
                                    <h3>Secured Output</h3>
                                </div>
                            """, unsafe_allow_html=True)
                            st.success(final_answer)
                            with st.expander("View Anonymized Pipeline"):
                                st.code(safe_text, language="text")

with tab_auditor:
    st.markdown("""
        <div class='section-header' style='border-bottom: none; margin-bottom: 2.5rem;'>
            <i class="fas fa-chart-line"></i>
            <h2 style='margin: 0; font-size: 1.875rem;'>Governance & Compliance Dashboard</h2>
        </div>
    """, unsafe_allow_html=True)
    
    if os.path.exists(AUDIT_FILE):
        with open(AUDIT_FILE, "r") as f:
            try:
                data = json.load(f)
                if data:
                    df = pd.DataFrame(data)
                    
                    st.markdown("""
                        <div style='margin-bottom: 1rem;'>
                            <h3 style='font-size: 1.125rem; color: #6b7280; font-weight: 600; text-transform: uppercase; letter-spacing: 0.5px;'>
                                Real-Time Intelligence
                            </h3>
                        </div>
                    """, unsafe_allow_html=True)
                    
                    m1, m2, m3 = st.columns(3)
                    
                    with m1:
                        st.markdown("""
                            <div class='metric-card'>
                                <i class="fas fa-server" style='color: #6366f1;'></i>
                            </div>
                        """, unsafe_allow_html=True)
                        st.metric("Total Operations", len(df))
                    
                    with m2:
                        st.markdown("""
                            <div class='metric-card'>
                                <i class="fas fa-shield-halved" style='color: #10b981;'></i>
                            </div>
                        """, unsafe_allow_html=True)
                        st.metric("PII Instances Blocked", df["blocked_items"].sum())
                    
                    with m3:
                        all_risks = []
                        for risks in df["risk_types"]:
                            all_risks.extend(risks)
                        most_common = max(set(all_risks), key=all_risks.count) if all_risks else "None"
                        st.markdown("""
                            <div class='metric-card'>
                                <i class="fas fa-triangle-exclamation" style='color: #f59e0b;'></i>
                            </div>
                        """, unsafe_allow_html=True)
                        st.metric("Primary Threat Vector", most_common)
                    
                    st.markdown("<br><br>", unsafe_allow_html=True)

                    st.markdown("""
                        <div style='margin-bottom: 1rem; margin-top: 2rem;'>
                            <h3 style='font-size: 1.125rem; color: #6b7280; font-weight: 600; text-transform: uppercase; letter-spacing: 0.5px;'>
                                Activity Analytics
                            </h3>
                        </div>
                    """, unsafe_allow_html=True)
                    
                    c1, c2 = st.columns(2, gap="large")
                    with c1:
                        st.markdown("""
                            <div style='display: flex; align-items: center; gap: 0.5rem; margin-bottom: 1rem;'>
                                <i class="fas fa-chart-area" style='color: #6366f1; font-size: 1rem;'></i>
                                <h4 style='margin: 0; font-size: 1rem; font-weight: 600; color: #374151;'>Protection Activity Timeline</h4>
                            </div>
                        """, unsafe_allow_html=True)
                        st.line_chart(df["blocked_items"], use_container_width=True)
                    
                    with c2:
                        st.markdown("""
                            <div style='display: flex; align-items: center; gap: 0.5rem; margin-bottom: 1rem;'>
                                <i class="fas fa-chart-column" style='color: #6366f1; font-size: 1rem;'></i>
                                <h4 style='margin: 0; font-size: 1rem; font-weight: 600; color: #374151;'>Threat Distribution Analysis</h4>
                            </div>
                        """, unsafe_allow_html=True)
                        if all_risks:
                            risk_counts = pd.Series(all_risks).value_counts()
                            st.bar_chart(risk_counts, use_container_width=True)
                        else:
                            st.markdown("""
                                <div class='custom-alert alert-info'>
                                    <i class="fas fa-info-circle"></i>
                                    <span>No threats detected in current dataset</span>
                                </div>
                            """, unsafe_allow_html=True)

                    st.markdown("<br><br>", unsafe_allow_html=True)

                    # Export Section
                    st.markdown("""
                        <div style='margin-bottom: 1.5rem; margin-top: 2rem;'>
                            <h3 style='font-size: 1.125rem; color: #6b7280; font-weight: 600; text-transform: uppercase; letter-spacing: 0.5px;'>
                                Compliance Reporting
                            </h3>
                        </div>
                    """, unsafe_allow_html=True)
                    
                    csv = df.to_csv(index=False).encode('utf-8')
                    
                    col_export1, col_export2 = st.columns([2, 1])
                    with col_export1:
                        st.download_button(
                            label="Export Audit Report (CSV)",
                            data=csv,
                            file_name=f"compliance_report_{datetime.datetime.now().strftime('%Y%m%d_%H%M%S')}.csv",
                            mime="text/csv",
                            use_container_width=True
                        )
                    
                    st.markdown("<br>", unsafe_allow_html=True)
                    st.dataframe(df, use_container_width=True, height=400)
                else:
                    st.markdown("""
                        <div class='custom-alert alert-info'>
                            <i class="fas fa-inbox"></i>
                            <span>Audit log initialized. Process data to generate reports.</span>
                        </div>
                    """, unsafe_allow_html=True)
            except Exception as e:
                st.markdown(f"""
                    <div class='custom-alert alert-error'>
                        <i class="fas fa-exclamation-triangle"></i>
                        <span>Error loading audit data: {e}</span>
                    </div>
                """, unsafe_allow_html=True)
    else:
        st.markdown("""
            <div class='custom-alert alert-info'>
                <i class="fas fa-inbox"></i>
                <span>No audit trail detected. Begin processing to establish compliance records.</span>
            </div>
        """, unsafe_allow_html=True)